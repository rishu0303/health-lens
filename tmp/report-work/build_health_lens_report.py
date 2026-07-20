from pathlib import Path
from textwrap import dedent

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/rishukumar/Desktop/Projects/health-lens")
WORK = ROOT / "tmp/report-work"
OUT_DIR = ROOT / "output/report"
OUT_DIR.mkdir(parents=True, exist_ok=True)
FINAL_DOCX = OUT_DIR / "Anjali_Rani_Health_Lens_Project_Report.docx"
LOGO = WORK / "reference-media/image1.png"
SCREENSHOTS = WORK / "screenshots"
DIAGRAMS = WORK / "diagrams"
DIAGRAMS.mkdir(parents=True, exist_ok=True)

BLUE = "0B3A67"
LIGHT_BLUE = "E8EEF5"
TEAL = "0F766E"
AMBER = "8B3F0A"
GRAY = "F4F6F9"
DARK = "1F2937"


def font_path(name="Arial.ttf"):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return None


def get_font(size, bold=False):
    path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else font_path()
    if path and Path(path).exists():
        return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded_rect(draw, xy, fill, outline="#CBD5E1", radius=16, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, p1, p2, color="#334155", width=4):
    draw.line([p1, p2], fill=color, width=width)
    x1, y1 = p1
    x2, y2 = p2
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - sign * 14, y2 - 8), (x2 - sign * 14, y2 + 8)]
    else:
        sign = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 8, y2 - sign * 14), (x2 + 8, y2 - sign * 14)]
    draw.polygon(pts, fill=color)


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines = []
    line = ""
    for word in words:
        test = f"{line} {word}".strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def as_hex(color):
    if isinstance(color, str) and color.startswith("#"):
        return color
    return f"#{color}"


def draw_box(draw, xy, title, subtitle="", fill="#FFFFFF", outline="#CBD5E1", title_color=DARK):
    x1, y1, x2, y2 = xy
    rounded_rect(draw, xy, fill, outline, radius=18, width=2)
    title_font = get_font(24, True)
    body_font = get_font(17)
    draw.text((x1 + 20, y1 + 18), title, fill=as_hex(title_color), font=title_font)
    y = y1 + 56
    for line in wrap_text(draw, subtitle, body_font, x2 - x1 - 40):
        draw.text((x1 + 20, y), line, fill="#475569", font=body_font)
        y += 24


def save_overall_flow():
    im = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 30), "Health Lens End-to-End Project Flow", fill=f"#{BLUE}", font=get_font(36, True))
    boxes = [
        ((60, 130, 330, 260), "User", "Logs in, uploads medical report, asks questions."),
        ((430, 130, 700, 260), "Frontend", "React/Vite dashboard, report history, upload and chat UI."),
        ((800, 130, 1070, 260), "Backend API", "Express routes, JWT auth, uploads, chat, reports."),
        ((1170, 130, 1440, 260), "Database", "MongoDB user, report, extracted fields and chat history."),
        ((430, 430, 700, 580), "Extraction", "PDF text/OCR is converted into structured parameters by Gemini + Zod."),
        ((800, 430, 1070, 580), "RAG Layer", "Knowledge PDFs are chunked, embedded, retrieved and cited."),
        ((1170, 430, 1440, 580), "Safe Answer", "Unified assistant separates report interpretation and education."),
    ]
    for xy, title, sub in boxes:
        fill = "#FFF7ED" if title in {"User", "Frontend"} else "#F0FDFA" if title in {"RAG Layer", "Safe Answer"} else "#F8FAFC"
        outline = "#C2410C" if title in {"User", "Frontend"} else "#0F766E" if title in {"RAG Layer", "Safe Answer"} else "#94A3B8"
        draw_box(d, xy, title, sub, fill=fill, outline=outline)
    arrow(d, (330, 195), (430, 195))
    arrow(d, (700, 195), (800, 195))
    arrow(d, (1070, 195), (1170, 195))
    arrow(d, (565, 260), (565, 430))
    arrow(d, (700, 505), (800, 505))
    arrow(d, (1070, 505), (1170, 505))
    arrow(d, (1305, 430), (1305, 260))
    out = DIAGRAMS / "overall-flow.png"
    im.save(out)
    return out


def save_upload_flow():
    im = Image.new("RGB", (1500, 620), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 30), "Report Upload and Analysis Flow", fill=f"#{BLUE}", font=get_font(34, True))
    labels = [
        ("PDF Upload", "Multer accepts PDF file and stores it temporarily."),
        ("Text Extraction", "pdf-parse extracts PDF text; image OCR support uses Tesseract."),
        ("AI Structuring", "Gemini returns reportType, summary, abnormal values and parameters."),
        ("Schema Guard", "Zod enforces parameter/value/range/status structure."),
        ("Persistence", "Report model stores extracted text, findings and suggested questions."),
    ]
    x = 55
    for i, (title, sub) in enumerate(labels):
        draw_box(d, (x, 150, x + 245, 330), title, sub, fill="#F8FAFC", outline="#64748B")
        if i < len(labels) - 1:
            arrow(d, (x + 245, 240), (x + 315, 240), color=f"#{AMBER}")
        x += 300
    out = DIAGRAMS / "upload-flow.png"
    im.save(out)
    return out


def save_rag_flow():
    im = Image.new("RGB", (1500, 680), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 30), "Knowledge Retrieval and Unified Chat Flow", fill=f"#{BLUE}", font=get_font(34, True))
    top = [
        ((80, 130, 380, 260), "Knowledge PDFs", "CBC, HbA1c, Thyroid, Vitamin D, Vitamin B12 and diabetes sources."),
        ((480, 130, 780, 260), "Chunking", "Recursive splitter creates overlapping chunks with source/page metadata."),
        ((880, 130, 1180, 260), "Embeddings", "Gemini embeddings or deterministic local fallback vectors."),
    ]
    bottom = [
        ((80, 420, 380, 550), "Question", "User asks an educational or report-specific question."),
        ((480, 420, 780, 550), "Top-K Retrieval", "Multi-query search, threshold filtering, deduplication and citations."),
        ((880, 420, 1180, 550), "Unified Answer", "Report values only from report; education only from retrieved context."),
    ]
    for xy, title, sub in top + bottom:
        draw_box(d, xy, title, sub, fill="#F0FDFA", outline="#0F766E")
    arrow(d, (380, 195), (480, 195), color="#0F766E")
    arrow(d, (780, 195), (880, 195), color="#0F766E")
    arrow(d, (230, 420), (230, 260), color="#64748B")
    arrow(d, (380, 485), (480, 485), color="#0F766E")
    arrow(d, (780, 485), (880, 485), color="#0F766E")
    draw_box(d, (1260, 270, 1440, 420), "Safety Layer", "Refuses diagnosis/prescription and adds urgent warnings.", fill="#FEF2F2", outline="#B91C1C", title_color="#991B1B")
    arrow(d, (1180, 485), (1320, 420), color="#B91C1C")
    out = DIAGRAMS / "rag-flow.png"
    im.save(out)
    return out


def save_eval_flow():
    im = Image.new("RGB", (1500, 720), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 30), "Evaluation Architecture", fill=f"#{BLUE}", font=get_font(34, True))
    stages = [
        ("Report Parsing", "Extraction Accuracy", "Correct fields / expected fields"),
        ("Retriever", "Hit Rate@5", "Relevant chunk in top five"),
        ("Generator", "Answer Correctness", "Manual score 0, 1 or 2"),
        ("Grounding", "Hallucination Rate", "Unsupported answers / reviewed answers"),
        ("Guardrails", "Safety Compliance", "Passed safety cases / total cases"),
    ]
    x = 45
    for i, (stage, metric, formula) in enumerate(stages):
        draw_box(d, (x, 165, x + 250, 360), stage, f"{metric}\n{formula}", fill="#F8FAFC", outline="#0B3A67", title_color=f"#{BLUE}")
        if i < len(stages) - 1:
            arrow(d, (x + 250, 260), (x + 285, 260), color="#0B3A67")
        x += 290
    results = [
        ("Extraction", "100% (28/28)", "Measured"),
        ("Retrieval", "60% (3/5)", "Measured"),
        ("Correctness", "Pending", "Manual review"),
        ("Hallucination", "Pending", "Manual review"),
        ("Safety", "100% (4/4)", "Measured"),
    ]
    d.text((60, 465), "Current prototype baseline", fill="#111827", font=get_font(26, True))
    for i, (name, score, status) in enumerate(results):
        x = 60 + i * 280
        rounded_rect(d, (x, 510, x + 240, 630), "#FFF7ED" if score == "Pending" else "#ECFDF5", "#F59E0B" if score == "Pending" else "#059669", radius=14, width=2)
        d.text((x + 18, 530), name, fill="#334155", font=get_font(18, True))
        d.text((x + 18, 558), score, fill="#111827", font=get_font(25, True))
        d.text((x + 18, 593), status, fill="#475569", font=get_font(16))
    out = DIAGRAMS / "evaluation-flow.png"
    im.save(out)
    return out


def save_model_diagram():
    im = Image.new("RGB", (1500, 660), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 30), "Data and Module Organization", fill=f"#{BLUE}", font=get_font(34, True))
    boxes = [
        ((80, 130, 420, 290), "User Model", "name, email, password hash, timestamps"),
        ((560, 130, 980, 310), "Report Model", "userId, originalFileName, extractedText, reportType, summary, abnormalValues, suggestions, suggestedQuestions, parameters, chatHistory"),
        ((1100, 130, 1420, 290), "Knowledge Cache", "chunks, vectors, source/page metadata and indexed timestamp"),
        ((320, 420, 640, 560), "Routes", "authRoutes, reportRoutes, chatRoutes"),
        ((860, 420, 1180, 560), "Services", "aiService, ragService, unifiedChatService, pdfService, ocrService"),
    ]
    for xy, title, sub in boxes:
        draw_box(d, xy, title, sub, fill="#F8FAFC", outline="#64748B")
    arrow(d, (420, 210), (560, 220), color="#64748B")
    arrow(d, (980, 220), (1100, 210), color="#64748B")
    arrow(d, (640, 490), (860, 490), color="#64748B")
    arrow(d, (720, 420), (760, 310), color="#64748B")
    out = DIAGRAMS / "data-modules.png"
    im.save(out)
    return out


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2EC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clear_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)


def set_table_widths(table, widths):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def set_run_font(run, size=None, bold=None, color=None, name="Lato"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, align=None, bold=False, size=None, color=None, after=6, before=0, line=1.12, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    pf.keep_with_next = keep
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level <= 2 else "1F4D78")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = add_para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, color="475569", after=8, before=2)
    return p


def add_image(doc, path, width=6.1, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    p.paragraph_format.space_after = Pt(2)
    if caption:
        add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        set_run_font(r, size=9.5, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
            p.paragraph_format.space_after = Pt(0)
    widths = widths or [int(9360 / len(headers))] * len(headers)
    set_table_geometry(table, widths)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
    return table


def set_document_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.45)
    sec.footer_distance = Inches(0.45)
    sec.different_first_page_header_footer = True
    styles = doc.styles
    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Lato"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Lato")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Lato")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Lato"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Lato")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Lato")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    footer = sec.footer
    p = footer.paragraphs[0]
    r = p.add_run("Health Lens Project Report | ")
    set_run_font(r, size=8, color="64748B")
    add_page_number(p)


def page_break(doc):
    doc.add_page_break()


def add_cover(doc):
    for _ in range(2):
        add_para(doc, "")
    add_para(doc, "HEALTH LENS: AI MEDICAL REPORT ANALYZER AND EVALUATION FRAMEWORK", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, after=22)
    add_para(doc, "A project report submitted in partial fulfilment of the requirements", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, after=2)
    add_para(doc, "for the research internship", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, after=20)
    add_para(doc, "MASTER OF COMPUTER APPLICATIONS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11, after=20)
    add_para(doc, "Submitted by:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10, after=6)
    add_para(doc, "ANJALI RANI   24204031205", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10, after=4)
    add_para(doc, "MANIT Bhopal | aranjalijangra@gmail.com", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, after=18)
    add_para(doc, "Submitted to:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10, after=6)
    add_para(doc, "Dr. Vinay Raj", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10, after=2)
    add_para(doc, "Assistant Professor, Department of Computer Applications", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, after=12)
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(0.95))
    add_para(doc, "DEPARTMENT OF COMPUTER APPLICATIONS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9.5, after=2)
    add_para(doc, "NATIONAL INSTITUTE OF TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9.5, after=2)
    add_para(doc, "TIRUCHIRAPPALLI - 620015", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9.5, after=8)
    add_para(doc, "SUMMER INTERNSHIP: MAY - JULY 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9.5, after=0)
    page_break(doc)


def add_front_matter(doc):
    add_para(doc, "BONAFIDE CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=18)
    add_para(
        doc,
        "This is to certify that the project entitled \"Health Lens: AI Medical Report Analyzer and Evaluation Framework\" is a project work successfully carried out by Anjali Rani (Roll No. 24204031205), student of MANIT Bhopal, during her distance research internship at the Department of Computer Applications, National Institute of Technology, Tiruchirappalli, under the guidance of Dr. Vinay Raj. The internship was offered for a two-month duration starting from 15 May 2026, as per the formal internship offer dated 30 April 2026.",
        size=10.5,
        after=10,
    )
    add_para(
        doc,
        "The work reported here includes the analysis, implementation, and evaluation of an AI-assisted medical-report understanding system with retrieval-augmented educational support and a five-metric evaluation framework.",
        size=10.5,
        after=28,
    )
    table = add_table(
        doc,
        ["Dr. Vinay Raj", "Head of the Department"],
        [["Mentor", "Department of Computer Applications"]],
        widths=[4680, 4680],
    )
    page_break(doc)

    add_para(doc, "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=14)
    for para in [
        "Health Lens is an AI-assisted medical-report analysis and educational question-answering prototype. The system allows an authenticated user to upload a medical report, extract readable text from the document, convert the extracted information into structured report fields, store the analysis in a patient report history, and ask contextual questions through a unified assistant.",
        "The implementation combines a React/Vite frontend, an Express and MongoDB backend, Gemini-based structured generation, PDF text extraction, OCR support, and a retrieval-augmented generation layer over a local medical knowledge base. The assistant separates report interpretation from general educational information, preserves patient-specific values from the uploaded report, and avoids unsafe diagnosis or prescription behaviour.",
        "A recent evaluation module was added to make the project measurable. It decomposes quality into five metrics: Report Extraction Accuracy, Retrieval Hit Rate@5, Answer Correctness, Hallucination Rate, and Safety Compliance Rate. The current prototype baseline reports 28/28 extraction field matches, 3/5 retrieval hits at K=5, and 4/4 safety cases passing, while answer correctness and hallucination remain intentionally human-reviewed. This report documents the system architecture, implementation, frontend workflow, evaluation methodology, limitations, and future research direction.",
    ]:
        add_para(doc, para, size=10.5, after=8)
    page_break(doc)

    add_para(doc, "ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=14)
    for para in [
        "I express my sincere gratitude to Dr. Vinay Raj, Assistant Professor, Department of Computer Applications, National Institute of Technology, Tiruchirappalli, for his guidance and support throughout this research internship. His direction helped shape the project into an end-to-end system with clear implementation and evaluation goals.",
        "I also thank the Department of Computer Applications, NIT Tiruchirappalli, for providing the academic setting for this internship. I am grateful to MANIT Bhopal for supporting my academic development, and to everyone who provided feedback during the project.",
    ]:
        add_para(doc, para, size=10.5, after=8)
    add_para(doc, "Anjali Rani", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=10.5, after=2)
    add_para(doc, "(24204031205)", align=WD_ALIGN_PARAGRAPH.RIGHT, size=10, after=0)
    page_break(doc)

    add_para(doc, "Table of contents", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, color=BLUE, after=12)
    toc = [
        ("BONAFIDE CERTIFICATE", "2", 0, True),
        ("ABSTRACT", "3", 0, True),
        ("ACKNOWLEDGEMENT", "4", 0, True),
        ("LIST OF ABBREVIATIONS", "6", 0, True),
        ("CHAPTER 1 - INTRODUCTION", "7", 0, True),
        ("1.1 Background", "7", 1, False),
        ("1.2 Problem Statement", "7", 1, False),
        ("1.3 Objectives", "7", 1, False),
        ("1.4 Scope and Contributions", "8", 1, False),
        ("CHAPTER 2 - TECHNICAL BACKGROUND", "9", 0, True),
        ("CHAPTER 3 - SYSTEM REQUIREMENTS AND ARCHITECTURE", "10", 0, True),
        ("CHAPTER 4 - IMPLEMENTATION", "12", 0, True),
        ("CHAPTER 5 - FRONTEND WORKFLOW", "15", 0, True),
        ("CHAPTER 6 - EVALUATION FRAMEWORK", "17", 0, True),
        ("CHAPTER 7 - RESULTS AND DISCUSSION", "19", 0, True),
        ("CHAPTER 8 - LIMITATIONS AND FUTURE ENHANCEMENTS", "20", 0, True),
        ("CHAPTER 9 - CONCLUSION", "21", 0, True),
        ("REFERENCES", "22", 0, True),
        ("APPENDIX: CODE MAP AND VERIFICATION", "23", 0, True),
    ]
    toc_table = doc.add_table(rows=0, cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for title, page, indent, is_major in toc:
        cells = toc_table.add_row().cells
        clear_cell_border(cells[0])
        clear_cell_border(cells[1])
        set_cell_margins(cells[0], top=25, bottom=25, start=0 if indent == 0 else 320, end=120)
        set_cell_margins(cells[1], top=25, bottom=25, start=120, end=0)
        title_p = cells[0].paragraphs[0]
        title_p.paragraph_format.space_after = Pt(0)
        title_run = title_p.add_run(title)
        set_run_font(title_run, size=10.2, bold=is_major)
        page_p = cells[1].paragraphs[0]
        page_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        page_p.paragraph_format.space_after = Pt(0)
        page_run = page_p.add_run(page)
        set_run_font(page_run, size=10.2, bold=is_major)
    set_table_widths(toc_table, [8200, 1160])
    page_break(doc)

    add_para(doc, "LIST OF ABBREVIATIONS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=10)
    add_table(
        doc,
        ["Abbreviation", "Expansion"],
        [
            ["AI", "Artificial Intelligence"],
            ["API", "Application Programming Interface"],
            ["CBC", "Complete Blood Count"],
            ["JWT", "JSON Web Token"],
            ["LLM", "Large Language Model"],
            ["OCR", "Optical Character Recognition"],
            ["PDF", "Portable Document Format"],
            ["RAG", "Retrieval-Augmented Generation"],
            ["UI", "User Interface"],
            ["Zod", "TypeScript/JavaScript schema validation library"],
        ],
        widths=[2300, 7060],
    )
    page_break(doc)


def add_chapter_1(doc):
    add_heading(doc, "CHAPTER 1 - INTRODUCTION", 1)
    add_para(doc, "Health Lens was developed as a research internship project to explore how an AI system can help users understand medical reports in a safer and more structured way. The system is not intended to diagnose or prescribe treatment. Its purpose is to extract report information, explain medical concepts in patient-friendly language, and provide a measurable foundation for further research.", size=10.5)
    add_heading(doc, "1.1 Background", 1)
    add_para(doc, "Medical reports are often dense, abbreviated, and difficult for non-specialist users to interpret. Patients may see laboratory values, reference intervals, flags such as High or Low, and clinical terms without knowing which values are report-specific and which explanations are general educational background.", size=10.5)
    add_para(doc, "A modern LLM can summarize and explain text, but medical use requires guardrails. Patient-specific claims must come from the uploaded report, educational claims should be grounded in reliable source material, and unsafe requests such as diagnosis or prescription must be refused. Health Lens addresses this by combining structured extraction, RAG, and deterministic safety checks.", size=10.5)
    add_heading(doc, "1.2 Problem Statement", 1)
    add_para(doc, "The problem is to design and evaluate a medical-report assistant that can process uploaded reports, store extracted information, answer questions using both report context and an indexed knowledge base, and expose measurable quality indicators for research review.", size=10.5)
    add_bullets(doc, [
        "Extract key medical parameters, values, reference ranges and status labels from uploaded PDF reports.",
        "Provide a report-history interface so users can switch assistant context between reports.",
        "Retrieve educational evidence from a medical PDF knowledge base before generating an answer.",
        "Separate report interpretation from educational explanation in the assistant response.",
        "Evaluate extraction, retrieval, answer quality, hallucination and safety using labelled datasets.",
    ])
    add_heading(doc, "1.3 Objectives", 1)
    add_bullets(doc, [
        "Build a usable frontend for authentication, upload, report history, search, knowledge-base status and chat.",
        "Implement backend routes for authentication, report upload, knowledge-base synchronization and unified chat.",
        "Use Gemini structured output with Zod validation for report analysis and chat response classification.",
        "Implement a RAG layer that supports local vector retrieval and Pinecone-backed retrieval configuration.",
        "Add a five-metric evaluation harness and document current prototype baselines with sample-size caveats.",
    ])
    add_heading(doc, "1.4 Scope and Contributions", 1)
    add_para(doc, "The project scope is an educational prototype for medical-report understanding. It demonstrates a full-stack AI workflow and a research-style evaluation framework, but it does not claim clinical validation or readiness for unsupervised medical deployment.", size=10.5)
    add_table(doc, ["Contribution", "Evidence in project"], [
        ["Full-stack prototype", "React dashboard, Express API, MongoDB models, authenticated routes and report history."],
        ["Structured analysis", "Gemini model output is constrained by a Zod schema for report fields and status labels."],
        ["RAG support", "Knowledge PDFs are chunked, embedded, cached or indexed, retrieved and cited."],
        ["Safety layer", "Unified prompts and deterministic urgent-symptom warning logic restrict unsafe outputs."],
        ["Evaluation framework", "Five metrics, JSON datasets, unit tests and live evaluation commands are implemented."],
    ], widths=[2600, 6760])


def add_chapter_2(doc):
    add_heading(doc, "CHAPTER 2 - TECHNICAL BACKGROUND", 1)
    add_para(doc, "This chapter summarizes the technical concepts that shaped the system: structured extraction, retrieval-augmented generation, schema validation, embeddings, and medical safety boundaries.", size=10.5)
    add_heading(doc, "2.1 Structured Medical Report Extraction", 1)
    add_para(doc, "Structured extraction converts raw report text into predictable fields. In Health Lens, the target structure includes report type, summary, abnormal values, suggestions, suggested questions and a parameter array. Each parameter stores the test name, observed value, reference range and status.", size=10.5)
    add_heading(doc, "2.2 Retrieval-Augmented Generation", 1)
    add_para(doc, "RAG reduces unsupported educational answers by retrieving relevant chunks before generation. Instead of asking the LLM to rely only on its pretraining, the backend searches the indexed medical knowledge base and passes the retrieved evidence to the assistant prompt.", size=10.5)
    add_heading(doc, "2.3 Embeddings and Similarity", 1)
    add_para(doc, "The project represents chunks and questions as dense embedding vectors. Similarity search ranks chunks by vector closeness. The local provider uses cosine similarity over cached vectors; the Pinecone path supports vector database search when configured.", size=10.5)
    add_heading(doc, "2.4 Medical Safety Boundary", 1)
    add_para(doc, "The assistant is designed for educational support, not clinical decision-making. It refuses diagnosis and prescription requests, separates patient-specific report interpretation from general information, recommends professional review, and adds an emergency warning for urgent symptom patterns.", size=10.5)
    add_table(doc, ["Concept", "Role in Health Lens"], [
        ["Zod schema", "Constrains LLM output into expected report and chat structures."],
        ["Temperature 0", "Improves repeatability during evaluation and structured generation."],
        ["JWT authentication", "Protects report upload, history and chat endpoints."],
        ["Top-K retrieval", "Selects the most relevant chunks for prompt context and citation."],
        ["Manual review", "Handles answer correctness and hallucination labels where exact matching is inadequate."],
    ], widths=[2600, 6760])


def add_chapter_3(doc, diagrams):
    add_heading(doc, "CHAPTER 3 - SYSTEM REQUIREMENTS AND ARCHITECTURE", 1)
    add_heading(doc, "3.1 Functional Requirements", 1)
    add_bullets(doc, [
        "Allow users to register, log in and access report features through token-protected API calls.",
        "Upload a medical PDF and process it into extracted text and structured analysis fields.",
        "Maintain a searchable history of uploaded reports and report-specific chat history.",
        "Index medical knowledge-base PDFs and report readiness status to the frontend.",
        "Answer user questions with optional report context and retrieved educational evidence.",
        "Provide evaluation commands for metric calculation and manual answer review.",
    ])
    add_heading(doc, "3.2 Non-Functional Requirements", 1)
    add_table(doc, ["Requirement", "Implementation response"], [
        ["Reproducibility", "Metric formulas are pure functions with deterministic unit tests."],
        ["Safety", "Prompts, structured medical context labels and deterministic emergency warnings are used."],
        ["Modularity", "Routes, controllers, services, models and evaluation modules are separated."],
        ["Extensibility", "RAG provider can be local or Pinecone; datasets can grow by adding JSON cases."],
        ["Usability", "Dashboard shows report counts, findings, status, selected report and assistant context."],
    ], widths=[2800, 6560])
    add_heading(doc, "3.3 Overall Architecture", 1)
    add_image(doc, diagrams["overall"], width=6.2, caption="Figure 3.1: End-to-end Health Lens project flow.")
    add_heading(doc, "3.4 Data and Module Organization", 1)
    add_para(doc, "The backend separates persistence, routing and AI services. User documents are stored in MongoDB through the Report model, while knowledge-base embeddings may be stored in a local cache or external vector index.", size=10.5)
    add_image(doc, diagrams["modules"], width=6.2, caption="Figure 3.2: Data and module organization used by the backend.")


def add_chapter_4(doc, diagrams):
    add_heading(doc, "CHAPTER 4 - IMPLEMENTATION", 1)
    add_heading(doc, "4.1 Technology Stack", 1)
    add_table(doc, ["Layer", "Technology"], [
        ["Frontend", "React 19, Vite, React Router, Axios, Lucide React icons, Tailwind-style utility classes"],
        ["Backend", "Node.js, Express 5, CommonJS modules, CORS, rate limiting middleware"],
        ["Database", "MongoDB through Mongoose User and Report schemas"],
        ["AI", "Gemini chat model through LangChain Google GenAI, structured output and prompt templates"],
        ["Document input", "pdf-parse for PDF text, Tesseract.js OCR support for images"],
        ["Retrieval", "LangChain PDFLoader, RecursiveCharacterTextSplitter, Gemini embeddings, local cache or Pinecone"],
        ["Evaluation", "Node test runner, JSON golden datasets, interactive manual review"],
    ], widths=[2300, 7060])
    add_heading(doc, "4.2 Backend API", 1)
    add_para(doc, "The API entry point initializes MongoDB and the RAG service, configures CORS for the local frontend origins, and mounts authentication, report and chat routes. Protected endpoints use JWT authentication, and upload/chat endpoints use rate limiting.", size=10.5)
    add_table(doc, ["Endpoint group", "Purpose"], [
        ["/api/auth", "Register and login users; return a token for authenticated calls."],
        ["/api/reports/upload", "Accept uploaded reports, extract text, analyze content and persist the report."],
        ["/api/reports", "Fetch report history and individual report details."],
        ["/api/reports/knowledge-base/status", "Expose indexed knowledge-base readiness and chunk count."],
        ["/api/reports/knowledge-base/sync", "Force re-indexing of the medical knowledge base."],
        ["/api/chat", "Run unified report-plus-knowledge chat and save report chat history when applicable."],
    ], widths=[3100, 6260])
    add_heading(doc, "4.3 Report Upload and Extraction", 1)
    add_image(doc, diagrams["upload"], width=6.2, caption="Figure 4.1: Upload, text extraction and structured analysis flow.")
    add_para(doc, "The upload controller accepts a PDF file, extracts text, sends it to the AI analysis service, stores the structured response in MongoDB, and removes the temporary upload file. This design prevents the upload directory from becoming the long-term storage of user documents.", size=10.5)
    add_heading(doc, "4.4 RAG and Unified Chat", 1)
    add_image(doc, diagrams["rag"], width=6.2, caption="Figure 4.2: Retrieval and unified chat flow.")
    add_para(doc, "The unified chat controller retrieves knowledge context for every question and conditionally adds report context when a reportId is present. The prompt instructs the model to use uploaded report fields only for patient-specific values and knowledge-base context only for educational explanation.", size=10.5)
    add_heading(doc, "4.5 Safety Implementation", 1)
    add_bullets(doc, [
        "Unsafe user text is treated as untrusted; prompt-injection attempts are explicitly rejected.",
        "Diagnosis, prescribing and doctor-role requests are classified as safety refusals.",
        "Urgent symptom patterns trigger a deterministic emergency warning outside the model.",
        "A consistent educational disclaimer is added for medical contexts.",
        "Follow-up questions guide the user toward safer next steps without pretending to be a clinician.",
    ])


def add_chapter_5(doc):
    add_heading(doc, "CHAPTER 5 - FRONTEND WORKFLOW", 1)
    add_para(doc, "The frontend is a single-page React application named MedInsight. It organizes the experience into a left navigation rail, dashboard/report/knowledge panels, a persistent assistant panel, and searchable report history. Axios interceptors attach the saved token to backend requests.", size=10.5)
    add_heading(doc, "5.1 Dashboard and Report Context", 1)
    add_para(doc, "The dashboard summarizes total reports, monthly reports, reports with findings, and knowledge-base readiness. The selected report card shows the latest report summary, key parameters, status labels and suggested questions.", size=10.5)
    if (SCREENSHOTS / "03-report-chat.png").exists():
        add_image(doc, SCREENSHOTS / "03-report-chat.png", width=6.2, caption="Figure 5.1: Patient report context and assistant response in the frontend.")
    add_heading(doc, "5.2 Knowledge Base Management", 1)
    add_para(doc, "The Knowledge Base screen combines the upload entry point with index status controls. The user can upload a report PDF, check readiness, refresh status and force a sync of the knowledge base.", size=10.5)
    if (SCREENSHOTS / "04-knowledge-base.png").exists():
        add_image(doc, SCREENSHOTS / "04-knowledge-base.png", width=6.2, caption="Figure 5.2: Knowledge-base upload and readiness status screen.")
    add_heading(doc, "5.3 Important UI States", 1)
    add_table(doc, ["State", "Frontend behaviour"], [
        ["Loading reports", "Shows a refresh spinner and keeps the report panel usable."],
        ["No reports", "Displays an empty state and directs the user to upload a medical PDF."],
        ["Upload in progress", "Shows a processing indicator and disables the file input."],
        ["Report selected", "Switches assistant placeholder and chat history to report-specific context."],
        ["Knowledge unavailable", "Displays readiness or last error so users know why chat may be delayed."],
    ], widths=[2600, 6760])


def add_chapter_6(doc, diagrams):
    add_heading(doc, "CHAPTER 6 - EVALUATION FRAMEWORK", 1)
    add_para(doc, "The recent evaluation work is one of the most important parts of the project. Instead of reporting one vague final score, Health Lens measures separate failure points in the pipeline. This makes the results actionable: extraction errors, retrieval misses, unsupported answers and safety failures can be diagnosed independently.", size=10.5)
    add_image(doc, diagrams["evaluation"], width=6.2, caption="Figure 6.1: Evaluation metrics mapped to pipeline stages.")
    add_heading(doc, "6.1 Golden Evaluation Data", 1)
    add_table(doc, ["Dataset", "Purpose"], [
        ["report-extraction.json", "Synthetic report text and expected structured parameter fields."],
        ["retrieval.json", "Questions with labelled relevant source/page/text chunk expectations."],
        ["answer-quality.json", "Questions and prepared reference answers for human review."],
        ["safety.json", "Diagnosis, prescription, prompt-injection and urgent-symptom cases."],
    ], widths=[3000, 6360])
    add_heading(doc, "6.2 Metric Definitions", 1)
    add_table(doc, ["Metric", "Formula / scoring", "Interpretation"], [
        ["Report Extraction Accuracy", "Correct fields / total expected fields", "Checks parameter, value, referenceRange and status for each expected parameter."],
        ["Retrieval Hit Rate@5", "Questions with relevant top-five chunk / total questions", "Measures whether the retriever supplies labelled evidence before generation."],
        ["Answer Correctness", "Average human score from 0 to 2", "Captures semantic answer quality using a prepared reference answer."],
        ["Hallucination Rate", "Unsupported answers / reviewed answers", "Counts answers containing at least one unsupported medical fact or patient value."],
        ["Safety Compliance Rate", "Safely handled cases / total safety cases", "Checks structured context, emergency warning and required refusal/urgency wording."],
    ], widths=[2500, 3200, 3660])
    add_heading(doc, "6.3 Evaluation Implementation", 1)
    add_para(doc, "The formulas are implemented in backend/evaluation/metrics.js. Deterministic unit tests in metrics.test.js validate edge cases such as missing parameters, relevant retrieval results ranked below K, pending manual labels and safety checks that require every expectation to pass.", size=10.5)
    add_para(doc, "runEvaluation.js executes live extraction, retrieval and safety cases against the application pipeline. reviewAnswerQuality.js generates answers, shows references and retrieved evidence, asks the reviewer for correctness and hallucination labels, and saves the review artifact to evaluation/results/answer-review.json.", size=10.5)
    add_heading(doc, "6.4 Current Baseline Results", 1)
    add_table(doc, ["Metric", "Current result", "Status", "Interpretation"], [
        ["Report Extraction Accuracy", "100% (28/28)", "Measured", "Clean synthetic smoke-test baseline across three report examples."],
        ["Retrieval Hit Rate@5", "60% (3/5)", "Measured", "Main improvement area; strict chunk-level evidence labels create useful misses."],
        ["Answer Correctness", "Pending review", "Manual", "Requires npm run evaluate:quality and human 0/1/2 labels."],
        ["Hallucination Rate", "Pending review", "Manual", "Requires reviewer judgement over generated claims and retrieved evidence."],
        ["Safety Compliance Rate", "100% (4/4)", "Measured", "All implemented diagnosis, prescription, prompt-injection and urgent-symptom cases passed."],
    ], widths=[2400, 1900, 1700, 3360])
    add_para(doc, "These percentages are prototype evaluation baselines, not clinical validation. The denominators are intentionally visible because a small 100% result should not be interpreted as universal safety or correctness.", size=10.5)


def add_chapter_7(doc):
    add_heading(doc, "CHAPTER 7 - RESULTS AND DISCUSSION", 1)
    add_heading(doc, "7.1 Verified Local Checks", 1)
    add_table(doc, ["Check", "Command", "Result"], [
        ["Metric unit tests", "npm test from backend", "7 tests passed; 0 failed."],
        ["Frontend production build", "npm run build from frontend", "Vite build completed; 106 modules transformed."],
        ["Reference report render", "render_docx.py on reference DOCX", "Reference pages rendered successfully for visual template inspection."],
    ], widths=[2400, 3700, 3260])
    add_heading(doc, "7.2 Interpretation of Evaluation Results", 1)
    add_para(doc, "The extraction result is strong for the current synthetic set because all expected fields matched after normalization. The most informative measured weakness is retrieval: three of five questions retrieved a strictly labelled relevant chunk in the top five. This points to chunking, threshold tuning, query expansion and reranking as the next experimental area.", size=10.5)
    add_para(doc, "The safety result is also promising within the current scope: the implemented diagnosis, prescription, prompt-injection and urgent-symptom cases passed. However, the suite is intentionally small, so the correct conclusion is that the implemented checks passed their current benchmark, not that all possible unsafe prompts are solved.", size=10.5)
    add_heading(doc, "7.3 Research Framing", 1)
    add_para(doc, "A useful research question emerging from this project is: How do chunking strategy, retrieval threshold, query expansion and reranking affect evidence retrieval, answer correctness, hallucination and safety in a medical-report RAG assistant?", size=10.5)
    add_table(doc, ["Variable type", "Examples"], [
        ["Independent variables", "Chunk size, overlap, top K, similarity threshold, multi-query search, reranker setting."],
        ["Dependent variables", "Hit Rate@5, correctness score, hallucination rate, safety compliance and latency."],
        ["Controls", "Same knowledge PDFs, same questions, same labels, temperature 0 and fixed model configuration."],
    ], widths=[2600, 6760])


def add_chapter_8(doc):
    add_heading(doc, "CHAPTER 8 - LIMITATIONS AND FUTURE ENHANCEMENTS", 1)
    add_heading(doc, "8.1 Limitations", 1)
    add_bullets(doc, [
        "The evaluation datasets are small and topic-specific, so reported percentages must include denominators.",
        "The extraction cases are synthetic and cleaner than many real scanned or photographed reports.",
        "Answer correctness and hallucination depend on reviewer judgement and should later use multiple reviewers.",
        "Retrieval labels are strict; a correct PDF/page can still count as a miss if the labelled evidence text is absent.",
        "The safety suite covers four high-value categories but not every adversarial prompt style or language.",
        "The system is educational and does not establish clinical validity, regulatory compliance or diagnostic accuracy.",
    ])
    add_heading(doc, "8.2 Future Enhancements", 1)
    add_bullets(doc, [
        "Expand the report-extraction dataset with de-identified real reports, OCR scans and varied laboratory formats.",
        "Tune chunk size, overlap, top K and similarity threshold using a development set before final testing.",
        "Add a reranker to improve evidence ordering after vector retrieval.",
        "Introduce claim-level hallucination labelling for deeper groundedness analysis.",
        "Add multilingual prompts and Indian clinical-report format variants to the safety and extraction suites.",
        "Develop a reviewer dashboard for manual correctness and hallucination labelling.",
    ])


def add_conclusion_refs_appendix(doc):
    add_heading(doc, "CHAPTER 9 - CONCLUSION", 1)
    add_para(doc, "Health Lens demonstrates a complete research-internship prototype for AI-assisted medical-report understanding. It includes a working full-stack application, structured report extraction, report-specific chat, RAG-based educational support, medical safety boundaries and an evaluation framework that separates the main sources of error.", size=10.5)
    add_para(doc, "The recently implemented evaluation suite is especially important because it turns the project from a feature demo into a measurable research artifact. The current baseline shows successful synthetic extraction and safety checks, while retrieval remains the clearest improvement target. The next phase should grow the labelled datasets and compare retrieval configurations under fixed controls.", size=10.5)
    page_break(doc)
    add_heading(doc, "REFERENCES", 1)
    refs = [
        "GitHub repository: https://github.com/aranjali1/health-lens.",
        "Project source code: /Users/rishukumar/Desktop/Projects/health-lens.",
        "Health Lens evaluation guide PDF supplied for this report: health-lens-evaluation-interview-guide.pdf.",
        "Internship offer letter supplied for this report: Anjali Rani.pdf, dated 30 April 2026.",
        "LangChain and Google Generative AI libraries used in backend services.",
        "MongoDB, Express, React, Node.js and Vite documentation for the MERN-style application stack.",
        "Tesseract.js and pdf-parse packages used for OCR and PDF text extraction support.",
    ]
    for ref in refs:
        add_para(doc, ref, size=10.2, after=4)
    page_break(doc)
    add_heading(doc, "APPENDIX: CODE MAP AND VERIFICATION", 1)
    add_table(doc, ["Path", "Role"], [
        ["https://github.com/aranjali1/health-lens", "Public GitHub repository for the Health Lens project."],
        ["frontend/src/App.jsx", "Dashboard, report history, upload panel, knowledge-base status and assistant UI."],
        ["frontend/src/Auth.jsx", "Login and registration screen."],
        ["backend/server.js", "Express app setup, CORS, routes and RAG initialization."],
        ["backend/controllers/reportController.js", "Upload, report history, knowledge-base status and legacy report chat handlers."],
        ["backend/controllers/chatController.js", "Unified report-plus-knowledge chat endpoint."],
        ["backend/services/aiService.js", "Structured medical report analysis with Gemini and Zod."],
        ["backend/services/ragService.js", "Knowledge-base loading, chunking, embedding, retrieval and citations."],
        ["backend/services/unifiedChatService.js", "Safety-aware unified answer generation."],
        ["backend/utils/medicalSafety.js", "Disclaimer and urgent symptom warning logic."],
        ["backend/evaluation/*.js", "Metric formulas, unit tests, live evaluation and manual review workflow."],
        ["backend/evaluation/data/*.json", "Golden labelled evaluation datasets."],
    ], widths=[3300, 6060])
    add_heading(doc, "Appendix A. Verification Summary", 2)
    add_bullets(doc, [
        "Backend metric tests: 7 passed, 0 failed.",
        "Frontend production build: completed successfully with 106 transformed modules.",
        "Reference report was rendered and inspected before creating this report.",
        "Final DOCX was rendered to page images for visual quality assurance.",
    ])


def build():
    diagrams = {
        "overall": save_overall_flow(),
        "upload": save_upload_flow(),
        "rag": save_rag_flow(),
        "evaluation": save_eval_flow(),
        "modules": save_model_diagram(),
    }
    doc = Document()
    set_document_styles(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_chapter_1(doc)
    page_break(doc)
    add_chapter_2(doc)
    page_break(doc)
    add_chapter_3(doc, diagrams)
    page_break(doc)
    add_chapter_4(doc, diagrams)
    page_break(doc)
    add_chapter_5(doc)
    page_break(doc)
    add_chapter_6(doc, diagrams)
    page_break(doc)
    add_chapter_7(doc)
    page_break(doc)
    add_chapter_8(doc)
    page_break(doc)
    add_conclusion_refs_appendix(doc)
    doc.save(FINAL_DOCX)
    print(FINAL_DOCX)


if __name__ == "__main__":
    build()
